import sys
from html import escape as html_escape
from typing import BinaryIO, Any
from .._base_converter import DocumentConverter, DocumentConverterResult
from .._exceptions import MissingDependencyException, MISSING_DEPENDENCY_MESSAGE
from .._stream_info import StreamInfo

# Try loading optional (but in this case, required) dependencies
# Save reporting of any exceptions for later
_xlsx_dependency_exc_info = None
try:
    import openpyxl
except ImportError:
    _xlsx_dependency_exc_info = sys.exc_info()

_xls_dependency_exc_info = None
try:
    import xlrd
except ImportError:
    _xls_dependency_exc_info = sys.exc_info()


def _rows_to_markdown(headers, rows):
    """Convert a list of header names and row data to a Markdown table string."""
    if not headers:
        return ""

    # Escape pipe characters and sanitize cell values
    def _cell(v):
        if v is None:
            return ""
        return html_escape(str(v)).replace("|", "\\|")

    lines = []
    lines.append("| " + " | ".join(_cell(h) for h in headers) + " |")
    lines.append("| " + " | ".join("---" for _ in headers) + " |")
    for row in rows:
        lines.append("| " + " | ".join(_cell(c) for c in row) + " |")
    return "\n".join(lines)

ACCEPTED_XLSX_MIME_TYPE_PREFIXES = [
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
]
ACCEPTED_XLSX_FILE_EXTENSIONS = [".xlsx"]

ACCEPTED_XLS_MIME_TYPE_PREFIXES = [
    "application/vnd.ms-excel",
    "application/excel",
]
ACCEPTED_XLS_FILE_EXTENSIONS = [".xls"]


class XlsxConverter(DocumentConverter):
    """
    Converts XLSX files to Markdown, with each sheet presented as a separate Markdown table.
    """

    def accepts(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,  # Options to pass to the converter
    ) -> bool:
        mimetype = (stream_info.mimetype or "").lower()
        extension = (stream_info.extension or "").lower()

        if extension in ACCEPTED_XLSX_FILE_EXTENSIONS:
            return True

        for prefix in ACCEPTED_XLSX_MIME_TYPE_PREFIXES:
            if mimetype.startswith(prefix):
                return True

        return False

    def convert(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,  # Options to pass to the converter
    ) -> DocumentConverterResult:
        # Check the dependencies
        if _xlsx_dependency_exc_info is not None:
            raise MissingDependencyException(
                MISSING_DEPENDENCY_MESSAGE.format(
                    converter=type(self).__name__,
                    extension=".xlsx",
                    feature="xlsx",
                )
            ) from _xlsx_dependency_exc_info[
                1
            ].with_traceback(  # type: ignore[union-attr]
                _xlsx_dependency_exc_info[2]
            )

        wb = openpyxl.load_workbook(file_stream, read_only=True, data_only=True)
        md_content = ""
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            all_rows = list(ws.iter_rows(values_only=True))
            if not all_rows:
                continue
            headers = [c if c is not None else "" for c in all_rows[0]]
            data = all_rows[1:]
            md_content += f"## {sheet_name}\n"
            md_content += _rows_to_markdown(headers, data) + "\n\n"
        wb.close()

        return DocumentConverterResult(markdown=md_content.strip())


class XlsConverter(DocumentConverter):
    """
    Converts XLS files to Markdown, with each sheet presented as a separate Markdown table.
    """

    def accepts(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,  # Options to pass to the converter
    ) -> bool:
        mimetype = (stream_info.mimetype or "").lower()
        extension = (stream_info.extension or "").lower()

        if extension in ACCEPTED_XLS_FILE_EXTENSIONS:
            return True

        for prefix in ACCEPTED_XLS_MIME_TYPE_PREFIXES:
            if mimetype.startswith(prefix):
                return True

        return False

    def convert(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,  # Options to pass to the converter
    ) -> DocumentConverterResult:
        # Load the dependencies
        if _xls_dependency_exc_info is not None:
            raise MissingDependencyException(
                MISSING_DEPENDENCY_MESSAGE.format(
                    converter=type(self).__name__,
                    extension=".xls",
                    feature="xls",
                )
            ) from _xls_dependency_exc_info[
                1
            ].with_traceback(  # type: ignore[union-attr]
                _xls_dependency_exc_info[2]
            )

        wb = xlrd.open_workbook(file_contents=file_stream.read())
        md_content = ""
        for sheet_name in wb.sheet_names():
            ws = wb.sheet_by_name(sheet_name)
            if ws.nrows == 0:
                continue
            headers = [ws.cell_value(0, c) for c in range(ws.ncols)]
            data = [
                [ws.cell_value(r, c) for c in range(ws.ncols)]
                for r in range(1, ws.nrows)
            ]
            md_content += f"## {sheet_name}\n"
            md_content += _rows_to_markdown(headers, data) + "\n\n"

        return DocumentConverterResult(markdown=md_content.strip())
